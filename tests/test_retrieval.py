"""检索模块测试"""

import asyncio
import pytest
from unittest.mock import AsyncMock, patch, MagicMock
from app.retrieval.retriever import Retriever
from app.retrieval.vector_store import (
    _metadata_matches_keyword_filter,
    _sync_keyword_search,
    expand_search_terms,
    normalize_search_text,
)
from app.retrieval.opensearch_store import (
    OpenSearchSearchResult,
    build_index_body,
    _document_source,
    _metadata_filters,
    _source_metadata,
)
from app.config import settings
from app.retrieval.ingest import (
    ingest_file,
    load_document,
    load_xlsx_document,
    sanitize_filename,
    save_uploaded_file,
    validate_file_extension,
)
from app.retrieval.document_registry import DocumentRegistry
from app.schemas import SourceItem


@pytest.fixture
def mock_retriever():
    retriever = Retriever()
    retriever._use_chroma = False
    retriever.refresh = MagicMock(return_value=False)
    return retriever


@pytest.fixture(autouse=True)
def mock_keyword_search_by_default():
    with patch("app.retrieval.vector_store.async_keyword_search", new=AsyncMock(return_value=[])):
        yield


@pytest.mark.asyncio
async def test_mock_search(mock_retriever):
    results = await mock_retriever.search("测试问题", top_k=3)
    assert len(results) == 3
    assert all(isinstance(r, SourceItem) for r in results)
    assert results[0].score > results[1].score


def test_validate_file_extension():
    assert validate_file_extension("test.pdf") is True
    assert validate_file_extension("test.txt") is True
    assert validate_file_extension("test.docx") is True
    assert validate_file_extension("test.xlsx") is True
    assert validate_file_extension("test.doc") is False
    assert validate_file_extension("test.docm") is False
    assert validate_file_extension("test.xls") is False


def test_sanitize_filename_removes_paths_and_invalid_chars():
    assert sanitize_filename("../evil.txt") == "evil.txt"
    assert sanitize_filename(r"C:\tmp\evil.txt") == "evil.txt"
    assert sanitize_filename("../evil.docx") == "evil.docx"
    assert sanitize_filename("../bad.xlsx") == "bad.xlsx"
    assert sanitize_filename('bad:name?.pdf') == "bad_name_.pdf"


def test_load_document_uses_docx_loader():
    with patch("langchain_community.document_loaders.Docx2txtLoader") as mock_loader:
        instance = mock_loader.return_value
        instance.load.return_value = ["doc"]

        result = load_document("test.docx")

    mock_loader.assert_called_once_with("test.docx")
    assert result == ["doc"]


def test_load_document_uses_xlsx_loader():
    with patch("app.retrieval.ingest.load_xlsx_document", return_value=["sheet doc"]) as mock_loader:
        result = load_document("test.xlsx")

    mock_loader.assert_called_once_with("test.xlsx")
    assert result == ["sheet doc"]


def test_load_xlsx_document_extracts_sheet_text(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SheetA"
    sheet.append(["名称", "数量"])
    sheet.append(["苹果", 3])
    workbook.save(path)

    docs = load_xlsx_document(str(path))

    assert len(docs) == 1
    assert "名称\t数量" in docs[0].page_content
    assert "苹果\t3" in docs[0].page_content
    assert docs[0].metadata["sheet"] == "SheetA"


def test_extract_query_terms_keeps_weather_focus_term():
    retriever = Retriever()

    terms = retriever._extract_query_terms("查找关于天气的历年毕业选题")

    assert "天气" in terms
    assert "历年" in terms
    assert "毕业选题" in terms


def test_extract_query_terms_expands_code_targets():
    retriever = Retriever()

    sort_terms = retriever._extract_query_terms("排序算法的python代码片段")
    sine_terms = retriever._extract_query_terms("绘制一条正弦曲线，曲线的样式是默认的")

    assert "快速排序" in sort_terms
    assert "quick_sort" in sort_terms
    assert "quicksort" in sort_terms
    assert "np.sin" in sine_terms
    assert "ax.plot" in sine_terms


def test_search_text_normalization_removes_function_words():
    assert normalize_search_text("帮我查找关于大语言模型的训练的内容") == "大语言模型训练"


def test_expand_search_terms_adds_llm_training_synonyms():
    terms = expand_search_terms(["大模型训练"])

    assert "大语言模型训练" in terms
    assert "大语言模型的训练" in terms
    assert "LLM" in terms


def test_extract_query_terms_expands_llm_training_targets():
    retriever = Retriever()

    terms = retriever._extract_query_terms("帮我查找关于大模型训练的内容")

    assert "大模型训练" in terms
    assert "大语言模型训练" in terms
    assert "大语言模型的训练" in terms


def test_merge_search_results_keeps_better_duplicate_score():
    retriever = Retriever()
    doc = MagicMock(
        page_content="以下是Python实现快速排序算法的示例代码：\ndef quick_sort(arr):\n    return arr",
        metadata={"source": "quick.txt", "document_id": "doc1", "chunk_index": 0},
    )

    results = retriever._merge_search_results([(doc, 1.5)], [(doc, -20)])

    assert results == [(doc, -20)]


def test_opensearch_index_body_uses_synonym_graph():
    body = build_index_body(use_ik_analyzer=False)

    filters = body["settings"]["analysis"]["filter"]
    content_mapping = body["mappings"]["properties"]["content"]
    mappings = body["mappings"]["properties"]

    assert filters["rag_synonyms"]["type"] == "synonym_graph"
    assert any("大模型" in synonym for synonym in filters["rag_synonyms"]["synonyms"])
    assert content_mapping["analyzer"] == "rag_index_analyzer"
    assert content_mapping["search_analyzer"] == "rag_search_analyzer"
    assert mappings["tenant_id"]["type"] == "keyword"
    assert mappings["owner_user_id"]["type"] == "keyword"
    assert mappings["chunk_id"]["type"] == "keyword"


def test_opensearch_personal_filter_requires_owner_and_personal_scope():
    filters = _metadata_filters({"owner_open_id": "ou_owner"})

    assert {"term": {"knowledge_base_type": "personal"}} in filters
    owner_filter = next(
        item for item in filters
        if {"term": {"owner_user_id": "ou_owner"}} in item.get("bool", {}).get("should", [])
    )
    assert {"term": {"owner_user_id": "ou_owner"}} in owner_filter["bool"]["should"]
    assert {"term": {"owner_open_id": "ou_owner"}} in owner_filter["bool"]["should"]


def test_opensearch_source_roundtrip_includes_tenant_and_owner_user_id():
    source = _document_source(
        "content",
        {
            "tenant_id": "tenant_a",
            "document_id": "doc1",
            "chunk_id": "doc1:0",
            "chunk_index": 0,
            "original_filename": "a.txt",
            "knowledge_base_type": "personal",
            "owner_user_id": "user_a",
            "owner_open_id": "user_a",
        },
    )
    metadata = _source_metadata(source)

    assert metadata["tenant_id"] == "tenant_a"
    assert metadata["owner_user_id"] == "user_a"
    assert metadata["owner_open_id"] == "user_a"
    assert metadata["chunk_id"] == "doc1:0"


@pytest.mark.asyncio
async def test_chroma_search_skips_local_fuzzy_when_opensearch_succeeds():
    from langchain_core.documents import Document

    retriever = Retriever()
    vector_doc = Document(
        page_content="向量命中内容",
        metadata={"source": "vector.txt", "knowledge_base_type": "enterprise", "document_id": "v1", "chunk_index": 0},
    )
    os_doc = Document(
        page_content="OpenSearch 命中内容",
        metadata={"source": "os.txt", "knowledge_base_type": "enterprise", "document_id": "o1", "chunk_index": 0},
    )

    with patch("app.retrieval.vector_store.async_similarity_search", new=AsyncMock(return_value=[(vector_doc, 0.2)])), \
         patch("app.retrieval.vector_store.async_get_document_chunks", new=AsyncMock(return_value=[])), \
         patch("app.retrieval.vector_store.async_keyword_search", new=AsyncMock(return_value=[])) as mock_keyword, \
         patch(
             "app.retrieval.opensearch_store.async_opensearch_search_with_status",
             new=AsyncMock(return_value=OpenSearchSearchResult(results=[(os_doc, 0.1)], available=True, succeeded=True)),
         ):
        results = await retriever._chroma_search("人工智能特点", top_k=5, knowledge_scope=["enterprise"])

    assert results
    mock_keyword.assert_not_awaited()


@pytest.mark.asyncio
async def test_chroma_search_skips_local_fuzzy_when_opensearch_succeeds_with_empty_results():
    from langchain_core.documents import Document

    retriever = Retriever()
    vector_doc = Document(
        page_content="向量命中内容",
        metadata={"source": "vector.txt", "knowledge_base_type": "enterprise", "document_id": "v1", "chunk_index": 0},
    )

    with patch("app.retrieval.vector_store.async_similarity_search", new=AsyncMock(return_value=[(vector_doc, 0.2)])), \
         patch("app.retrieval.vector_store.async_get_document_chunks", new=AsyncMock(return_value=[])), \
         patch("app.retrieval.vector_store.async_keyword_search", new=AsyncMock(return_value=[])) as mock_keyword, \
         patch(
             "app.retrieval.opensearch_store.async_opensearch_search_with_status",
             new=AsyncMock(return_value=OpenSearchSearchResult(results=[], available=True, succeeded=True)),
         ):
        results = await retriever._chroma_search("人工智能特点", top_k=5, knowledge_scope=["enterprise"])

    assert results
    mock_keyword.assert_not_awaited()


@pytest.mark.asyncio
async def test_chroma_search_uses_local_fuzzy_when_opensearch_unavailable():
    from langchain_core.documents import Document

    retriever = Retriever()
    vector_doc = Document(
        page_content="向量命中内容",
        metadata={"source": "vector.txt", "knowledge_base_type": "enterprise", "document_id": "v1", "chunk_index": 0},
    )
    keyword_doc = Document(
        page_content="本地 fuzzy 命中内容",
        metadata={"source": "keyword.txt", "knowledge_base_type": "enterprise", "document_id": "k1", "chunk_index": 0},
    )

    with patch("app.retrieval.vector_store.async_similarity_search", new=AsyncMock(return_value=[(vector_doc, 0.2)])), \
         patch("app.retrieval.vector_store.async_get_document_chunks", new=AsyncMock(return_value=[])), \
         patch("app.retrieval.vector_store.async_keyword_search", new=AsyncMock(return_value=[(keyword_doc, -1.0)])) as mock_keyword, \
         patch(
             "app.retrieval.opensearch_store.async_opensearch_search_with_status",
             new=AsyncMock(return_value=OpenSearchSearchResult(results=[], available=False, succeeded=False, error="unavailable")),
         ):
        results = await retriever._chroma_search("人工智能特点", top_k=5, knowledge_scope=["enterprise"])

    assert results
    mock_keyword.assert_awaited_once()


def test_keyword_search_matches_llm_training_without_particle():
    class FakeVectorStore:
        def get(self, include=None, **kwargs):
            return {
                "documents": ["课程主题：大语言模型的训练，包括预训练和微调。"],
                "metadatas": [{"source": "llm.txt", "knowledge_base_type": "enterprise"}],
            }

    with patch("app.retrieval.vector_store.get_vector_store", return_value=FakeVectorStore()):
        results = _sync_keyword_search(["大语言模型训练"], k=5, metadata_filter={"knowledge_base_type": "enterprise"})

    assert len(results) == 1
    assert "大语言模型的训练" in results[0][0].page_content


def test_keyword_search_matches_big_model_training_synonym():
    class FakeVectorStore:
        def get(self, include=None, **kwargs):
            return {
                "documents": ["课程主题：大语言模型的训练，包括预训练和微调。"],
                "metadatas": [{"source": "llm.txt", "knowledge_base_type": "enterprise"}],
            }

    with patch("app.retrieval.vector_store.get_vector_store", return_value=FakeVectorStore()):
        results = _sync_keyword_search(["帮我查找关于大模型训练的内容"], k=5, metadata_filter={"knowledge_base_type": "enterprise"})

    assert len(results) == 1
    assert "大语言模型的训练" in results[0][0].page_content


def test_keyword_search_does_not_overmatch_unrelated_model_phrases():
    class FakeVectorStore:
        def get(self, include=None, **kwargs):
            return {
                "documents": ["课程主题：大语言模型的训练，包括预训练和微调。"],
                "metadatas": [{"source": "llm.txt", "knowledge_base_type": "enterprise"}],
            }

    with patch("app.retrieval.vector_store.get_vector_store", return_value=FakeVectorStore()):
        car_results = _sync_keyword_search(["模型车训练"], k=5, metadata_filter={"knowledge_base_type": "enterprise"})
        exhibition_results = _sync_keyword_search(["大型模型展览"], k=5, metadata_filter={"knowledge_base_type": "enterprise"})

    assert car_results == []
    assert exhibition_results == []


def test_keyword_search_keeps_personal_owner_filter_for_fuzzy_match():
    class FakeVectorStore:
        def get(self, include=None, **kwargs):
            assert kwargs["where"] == {"owner_open_id": "ou_owner"}
            return {
                "documents": ["课程主题：大语言模型的训练。", "课程主题：大语言模型的训练。"],
                "metadatas": [
                    {"source": "mine.txt", "knowledge_base_type": "personal", "owner_open_id": "ou_owner"},
                    {"source": "other.txt", "knowledge_base_type": "personal", "owner_open_id": "ou_other"},
                ],
            }

    with patch("app.retrieval.vector_store.get_vector_store", return_value=FakeVectorStore()):
        results = _sync_keyword_search(["大模型训练"], k=5, metadata_filter={"owner_open_id": "ou_owner"})

    assert len(results) == 1
    assert results[0][0].metadata["source"] == "mine.txt"


def test_keyword_filter_accepts_owner_user_id_and_legacy_owner_open_id():
    metadata_filter = {"owner_user_id": "ou_owner", "tenant_id": settings.default_tenant_id}

    assert _metadata_matches_keyword_filter(
        {
            "knowledge_base_type": "personal",
            "owner_user_id": "ou_owner",
            "tenant_id": settings.default_tenant_id,
        },
        metadata_filter,
    )
    assert _metadata_matches_keyword_filter(
        {
            "knowledge_base_type": "personal",
            "owner_open_id": "ou_owner",
        },
        metadata_filter,
    )
    assert not _metadata_matches_keyword_filter(
        {
            "knowledge_base_type": "personal",
            "owner_user_id": "ou_other",
            "tenant_id": settings.default_tenant_id,
        },
        metadata_filter,
    )


def test_keyword_search_treats_legacy_enterprise_metadata_as_enterprise():
    class FakeVectorStore:
        def get(self, include=None, **kwargs):
            assert "where" not in kwargs
            return {
                "documents": ["课程主题：大语言模型的训练。"],
                "metadatas": [{"source": "legacy.txt"}],
            }

    with patch("app.retrieval.vector_store.get_vector_store", return_value=FakeVectorStore()):
        results = _sync_keyword_search(["大模型训练"], k=5, metadata_filter={"knowledge_base_type": "enterprise"})

    assert len(results) == 1
    assert results[0][0].metadata["source"] == "legacy.txt"


def test_save_uploaded_file_stays_inside_upload_dir(tmp_path):
    old_upload_dir = settings.upload_dir
    settings.upload_dir = str(tmp_path)
    try:
        for filename in ["../evil.txt", r"C:\tmp\evil.txt", "bad:name?.txt"]:
            saved_path = save_uploaded_file(b"hello", filename)
            resolved = __import__("pathlib").Path(saved_path).resolve()
            assert tmp_path.resolve() in resolved.parents
            assert ".." not in resolved.name
            assert ":" not in resolved.name
    finally:
        settings.upload_dir = old_upload_dir


def test_save_uploaded_file_uses_per_user_personal_path(tmp_path):
    saved_path = save_uploaded_file(
        b"hello",
        "doc.txt",
        upload_dir=str(tmp_path),
        tenant_id="tenant/a",
        owner_user_id="ou:test",
        knowledge_base_type="personal",
        document_id="doc123",
    )
    resolved = __import__("pathlib").Path(saved_path).resolve()

    assert tmp_path.resolve() in resolved.parents
    assert "personal" in resolved.parts
    assert "tenant_a" in resolved.parts
    assert "ou_test" in resolved.parts
    assert resolved.name.startswith("doc123_")


@patch("app.retrieval.vector_store.add_documents")
@patch("app.retrieval.ingest.split_documents")
@patch("app.retrieval.ingest.load_document")
def test_ingest_file(mock_load, mock_split, mock_add, tmp_path):
    mock_load.return_value = [MagicMock()]
    chunk1 = MagicMock(metadata={"page": 1})
    chunk2 = MagicMock(metadata={"page": 2})
    mock_split.return_value = [chunk1, chunk2]
    mock_add.return_value = 2

    with patch("app.retrieval.document_registry.document_registry", DocumentRegistry(str(tmp_path / "registry.sqlite3"))):
        result = ingest_file("/tmp/stored.pdf", original_filename="../test.pdf")
    assert result["chunks"] == 2
    assert "document_id" in result
    assert "filename" in result
    assert result["filename"] == "test.pdf"
    assert chunk1.metadata["document_id"] == result["document_id"]
    assert chunk1.metadata["tenant_id"] == settings.default_tenant_id
    assert chunk1.metadata["chunk_id"] == f"{result['document_id']}:0"
    assert chunk2.metadata["document_id"] == result["document_id"]
    assert chunk1.metadata["chunk_index"] == 0
    assert chunk2.metadata["chunk_index"] == 1
    assert chunk1.metadata["original_filename"] == "test.pdf"
    assert chunk1.metadata["stored_filename"] == "stored.pdf"
    assert chunk1.metadata["knowledge_base_type"] == "enterprise"
    assert chunk1.metadata["owner_user_id"] == ""
    assert chunk1.metadata["owner_open_id"] == ""
    assert chunk1.metadata["chat_id"] == ""
    assert chunk1.metadata["channel"] == "api"


@patch("app.retrieval.vector_store.add_documents")
@patch("app.retrieval.ingest.split_documents")
@patch("app.retrieval.ingest.load_document")
def test_ingest_file_personal_metadata(mock_load, mock_split, mock_add, tmp_path):
    mock_load.return_value = [MagicMock()]
    chunk = MagicMock(metadata={})
    mock_split.return_value = [chunk]
    mock_add.return_value = 1

    with patch("app.retrieval.document_registry.document_registry", DocumentRegistry(str(tmp_path / "registry.sqlite3"))):
        result = ingest_file(
            "/tmp/stored.txt",
            original_filename="我的资料.txt",
            knowledge_base_type="personal",
            owner_open_id="ou_test123",
            owner_user_id="ou_test123",
            chat_id="oc_test789",
            channel="feishu",
        )

    assert result["chunks"] == 1
    assert chunk.metadata["knowledge_base_type"] == "personal"
    assert chunk.metadata["tenant_id"] == settings.default_tenant_id
    assert chunk.metadata["owner_user_id"] == "ou_test123"
    assert chunk.metadata["owner_open_id"] == "ou_test123"
    assert chunk.metadata["chat_id"] == "oc_test789"
    assert chunk.metadata["channel"] == "feishu"


@patch("app.retrieval.opensearch_store.index_documents")
@patch("app.retrieval.vector_store.add_documents")
@patch("app.retrieval.ingest.split_documents")
@patch("app.retrieval.ingest.load_document")
def test_ingest_file_updates_opensearch_index(mock_load, mock_split, mock_add, mock_index, tmp_path):
    mock_load.return_value = [MagicMock()]
    chunk = MagicMock(metadata={})
    mock_split.return_value = [chunk]
    mock_add.return_value = 1
    mock_index.return_value = 1

    with patch("app.retrieval.document_registry.document_registry", DocumentRegistry(str(tmp_path / "registry.sqlite3"))):
        result = ingest_file("/tmp/stored.txt", original_filename="stored.txt")

    assert result["chunks"] == 1
    mock_index.assert_called_once_with([chunk])


@pytest.mark.asyncio
async def test_chroma_search_fallback_to_mock():
    retriever = Retriever()
    retriever._use_chroma = False
    retriever.refresh = MagicMock(return_value=False)
    results = await retriever.search("测试", top_k=2)
    assert len(results) > 0


def test_normalize_distance_score():
    retriever = Retriever()
    assert retriever._normalize_distance_score(0) == 1.0
    assert 0.0 < retriever._normalize_distance_score(1.5) < 1.0
    assert retriever._normalize_distance_score(None) == 0.0
    assert 0.0 <= retriever._normalize_distance_score(-1) <= 1.0


@pytest.mark.asyncio
async def test_chroma_exception_falls_back_to_mock():
    retriever = Retriever()
    retriever._use_chroma = True
    with patch.object(retriever, "_chroma_search", new=AsyncMock(side_effect=RuntimeError("boom"))):
        results = await retriever.search("测试", top_k=2)

    assert len(results) == 2
    assert all(item.source_type == "knowledge_base" for item in results)


@pytest.mark.asyncio
async def test_chroma_search_keeps_full_content_for_llm_and_short_snippet():
    retriever = Retriever()
    long_content = "开头" * 120 + "先小康后大同 阶级特征 高度的生产力"
    doc = MagicMock(
        page_content=long_content,
        metadata={"source": "4.pdf", "document_id": "doc1", "chunk_index": 0},
    )

    with patch(
        "app.retrieval.vector_store.async_similarity_search",
        new=AsyncMock(return_value=[(doc, 0)]),
    ), patch(
        "app.retrieval.vector_store.async_get_document_chunks",
        new=AsyncMock(return_value=[]),
    ):
        results = await retriever._chroma_search("历史影响", top_k=1)

    assert len(results) == 1
    assert results[0].snippet == long_content[:settings.rag_display_snippet_chars]
    assert "先小康后大同" in results[0].content
    assert "content" not in results[0].model_dump()


@pytest.mark.asyncio
async def test_chroma_search_expands_neighbor_chunks():
    retriever = Retriever()
    doc = MagicMock(
        page_content="5. 儒家大同思想的历史影响",
        metadata={"source": "4.pdf", "document_id": "doc1", "chunk_index": 1},
    )
    chunks = [
        {"content": "上一段", "metadata": {"chunk_index": 0}},
        {"content": "5. 儒家大同思想的历史影响", "metadata": {"chunk_index": 1}},
        {"content": "进步性、阶级特征、缺乏可实现途径", "metadata": {"chunk_index": 2}},
    ]

    with patch(
        "app.retrieval.vector_store.async_similarity_search",
        new=AsyncMock(return_value=[(doc, 0)]),
    ), patch(
        "app.retrieval.vector_store.async_get_document_chunks",
        new=AsyncMock(return_value=chunks),
    ):
        results = await retriever._chroma_search("历史影响", top_k=1)

    assert "上一段" in results[0].content
    assert "进步性、阶级特征、缺乏可实现途径" in results[0].content


@pytest.mark.asyncio
async def test_chroma_search_expands_more_chunks_for_global_questions():
    retriever = Retriever()
    doc = MagicMock(
        page_content="儒家大同思想的基本内容主要包括以下五个方面。",
        metadata={"source": "4.pdf", "document_id": "doc1", "chunk_index": 5},
    )
    chunks = [
        {"content": "社会制度：全民公有", "metadata": {"chunk_index": 5}},
        {"content": "管理制度：选贤与能", "metadata": {"chunk_index": 6}},
        {"content": "人际关系：讲信修睦", "metadata": {"chunk_index": 7}},
        {"content": "社会保障：人人得其所", "metadata": {"chunk_index": 8}},
        {"content": "劳动态度：各尽其力", "metadata": {"chunk_index": 9}},
    ]

    with patch(
        "app.retrieval.vector_store.async_similarity_search",
        new=AsyncMock(return_value=[(doc, 0)]),
    ), patch(
        "app.retrieval.vector_store.async_get_document_chunks",
        new=AsyncMock(return_value=chunks),
    ):
        results = await retriever._chroma_search("儒家大同思想的主要内容", top_k=1)

    assert "全民公有" in results[0].content
    assert "选贤与能" in results[0].content
    assert "讲信修睦" in results[0].content
    assert "人人得其所" in results[0].content
    assert "各尽其力" in results[0].content


@pytest.mark.asyncio
async def test_chroma_search_neighbor_failure_uses_original_chunk():
    retriever = Retriever()
    doc = MagicMock(
        page_content="原始命中切块",
        metadata={"source": "4.pdf", "document_id": "doc1", "chunk_index": 1},
    )

    with patch(
        "app.retrieval.vector_store.async_similarity_search",
        new=AsyncMock(return_value=[(doc, 0)]),
    ), patch(
        "app.retrieval.vector_store.async_get_document_chunks",
        new=AsyncMock(side_effect=RuntimeError("boom")),
    ):
        results = await retriever._chroma_search("历史影响", top_k=1)

    assert results[0].content == "原始命中切块"


@pytest.mark.asyncio
async def test_enterprise_scope_excludes_personal_sources():
    retriever = Retriever()
    enterprise_doc = MagicMock(
        page_content="企业知识",
        metadata={"source": "enterprise.txt", "knowledge_base_type": "enterprise"},
    )
    legacy_doc = MagicMock(
        page_content="旧企业知识",
        metadata={"source": "legacy.txt"},
    )
    personal_doc = MagicMock(
        page_content="个人知识",
        metadata={"source": "personal.txt", "knowledge_base_type": "personal", "owner_open_id": "ou_test123"},
    )

    with patch(
        "app.retrieval.vector_store.async_similarity_search",
        new=AsyncMock(return_value=[(enterprise_doc, 0), (legacy_doc, 0.1), (personal_doc, 0.2)]),
    ), patch(
        "app.retrieval.vector_store.async_get_document_chunks",
        new=AsyncMock(return_value=[]),
    ):
        results = await retriever._chroma_search("测试", top_k=5, knowledge_scope=["enterprise"])

    titles = [item.title for item in results]
    assert "enterprise.txt" in titles
    assert "legacy.txt" in titles
    assert "personal.txt" not in titles


@pytest.mark.asyncio
async def test_personal_scope_only_returns_matching_owner():
    retriever = Retriever()
    owner_doc = MagicMock(
        page_content="当前用户个人知识",
        metadata={"source": "mine.txt", "knowledge_base_type": "personal", "owner_open_id": "ou_test123"},
    )
    other_doc = MagicMock(
        page_content="其他用户个人知识",
        metadata={"source": "other.txt", "knowledge_base_type": "personal", "owner_open_id": "ou_other"},
    )
    enterprise_doc = MagicMock(
        page_content="企业知识",
        metadata={"source": "enterprise.txt", "knowledge_base_type": "enterprise"},
    )

    with patch(
        "app.retrieval.vector_store.async_similarity_search",
        new=AsyncMock(return_value=[(owner_doc, 0), (other_doc, 0.1), (enterprise_doc, 0.2)]),
    ), patch(
        "app.retrieval.vector_store.async_get_document_chunks",
        new=AsyncMock(return_value=[]),
    ):
        results = await retriever._chroma_search(
            "测试",
            top_k=5,
            knowledge_scope=["personal"],
            owner_open_id="ou_test123",
        )

    assert [item.title for item in results] == ["mine.txt"]


@pytest.mark.asyncio
async def test_personal_scope_falls_back_when_owner_filter_misses():
    retriever = Retriever()
    personal_doc = MagicMock(
        page_content="Python实现快速排序算法\ndef quick_sort(arr):\n    return arr",
        metadata={"source": "personal_quick_sort.txt", "knowledge_base_type": "personal", "owner_open_id": "ou_ingest_owner"},
    )

    async def fake_search(_query, k=5, metadata_filter=None):
        if metadata_filter == {"owner_open_id": "ou_request_user"}:
            return []
        if metadata_filter == {"knowledge_base_type": "personal"}:
            return [(personal_doc, 0)]
        return []

    with patch.object(settings, "personal_kb_strict_owner_filter", False), patch(
        "app.retrieval.vector_store.async_similarity_search",
        new=AsyncMock(side_effect=fake_search),
    ), patch(
        "app.retrieval.vector_store.async_get_document_chunks",
        new=AsyncMock(return_value=[]),
    ):
        results = await retriever._chroma_search(
            "Python实现快速排序算法",
            top_k=5,
            knowledge_scope=["personal"],
            owner_open_id="ou_request_user",
        )

    assert [item.title for item in results] == ["personal_quick_sort.txt"]


@pytest.mark.asyncio
async def test_personal_scope_reranks_exact_code_hit_before_vector_noise():
    retriever = Retriever()
    noise_docs = [
        MagicMock(
            page_content=f"个人知识库无关内容 {index}",
            metadata={"source": f"noise_{index}.txt", "knowledge_base_type": "personal", "owner_open_id": "ou_ingest_owner"},
        )
        for index in range(3)
    ]
    quick_sort_doc = MagicMock(
        page_content="以下是Python实现快速排序算法的示例代码：\ndef quick_sort(arr):\n    return arr",
        metadata={"source": "personal_quick_sort.txt", "knowledge_base_type": "personal", "owner_open_id": "ou_ingest_owner"},
    )

    async def fake_search(_query, k=5, metadata_filter=None):
        if metadata_filter == {"owner_open_id": "ou_request_user"}:
            return []
        if metadata_filter == {"knowledge_base_type": "personal"}:
            return [(doc, index * 0.01) for index, doc in enumerate(noise_docs)] + [(quick_sort_doc, 1.5)]
        return []

    with patch.object(settings, "personal_kb_strict_owner_filter", False), patch(
        "app.retrieval.vector_store.async_similarity_search",
        new=AsyncMock(side_effect=fake_search),
    ), patch(
        "app.retrieval.vector_store.async_get_document_chunks",
        new=AsyncMock(return_value=[]),
    ):
        results = await retriever._chroma_search(
            "个人知识库检索，找出“Python实现快速排序算法”的python代码片段",
            top_k=3,
            knowledge_scope=["personal"],
            owner_open_id="ou_request_user",
        )

    assert results[0].title == "personal_quick_sort.txt"


@pytest.mark.asyncio
async def test_personal_sort_algorithm_returns_quick_sort_before_unrelated_chunks():
    retriever = Retriever()
    noise_doc = MagicMock(
        page_content="用户：安装复杂吗？需要专业人员吗？\n客服：基础安装非常简单。",
        metadata={"source": "personal_noise.txt", "knowledge_base_type": "personal", "owner_open_id": "ou_test123"},
    )
    quick_sort_doc = MagicMock(
        page_content="以下是Python实现快速排序算法的示例代码：\n\ndef quick_sort(arr):\n    if len(arr) <= 1:\n        return arr",
        metadata={"source": "personal_quick_sort.txt", "knowledge_base_type": "personal", "owner_open_id": "ou_test123"},
    )

    async def fake_similarity_search(_query, k=5, metadata_filter=None):
        return [(noise_doc, 0), (quick_sort_doc, 1.5)]

    async def fake_keyword_search(terms, k=5, metadata_filter=None):
        assert "quick_sort" in terms
        return [(quick_sort_doc, -20)]

    with patch(
        "app.retrieval.vector_store.async_similarity_search",
        new=AsyncMock(side_effect=fake_similarity_search),
    ), patch(
        "app.retrieval.vector_store.async_keyword_search",
        new=AsyncMock(side_effect=fake_keyword_search),
    ), patch(
        "app.retrieval.vector_store.async_get_document_chunks",
        new=AsyncMock(return_value=[]),
    ):
        results = await retriever._chroma_search(
            "“排序算法”的python代码片段",
            top_k=3,
            knowledge_scope=["personal"],
            owner_open_id="ou_test123",
        )

    assert results[0].title == "personal_quick_sort.txt"
    assert "def quick_sort" in results[0].snippet


@pytest.mark.asyncio
async def test_personal_scope_strict_owner_filter_disables_fallback():
    retriever = Retriever()
    personal_doc = MagicMock(
        page_content="Python实现快速排序算法",
        metadata={"source": "personal_quick_sort.txt", "knowledge_base_type": "personal", "owner_open_id": "ou_ingest_owner"},
    )

    async def fake_search(_query, k=5, metadata_filter=None):
        if metadata_filter == {"owner_open_id": "ou_request_user"}:
            return []
        if metadata_filter == {"knowledge_base_type": "personal"}:
            return [(personal_doc, 0)]
        return []

    with patch.object(settings, "personal_kb_strict_owner_filter", True), patch(
        "app.retrieval.vector_store.async_similarity_search",
        new=AsyncMock(side_effect=fake_search),
    ), patch(
        "app.retrieval.vector_store.async_get_document_chunks",
        new=AsyncMock(return_value=[]),
    ):
        results = await retriever._chroma_search(
            "Python实现快速排序算法",
            top_k=5,
            knowledge_scope=["personal"],
            owner_open_id="ou_request_user",
        )

    assert results == []


@pytest.mark.asyncio
async def test_combined_scope_merges_enterprise_and_current_owner_personal():
    retriever = Retriever()
    enterprise_doc = MagicMock(
        page_content="企业知识",
        metadata={"source": "enterprise.txt", "knowledge_base_type": "enterprise"},
    )
    owner_doc = MagicMock(
        page_content="当前用户个人知识",
        metadata={"source": "mine.txt", "knowledge_base_type": "personal", "owner_open_id": "ou_test123"},
    )
    other_doc = MagicMock(
        page_content="其他用户个人知识",
        metadata={"source": "other.txt", "knowledge_base_type": "personal", "owner_open_id": "ou_other"},
    )

    async def fake_search(_query, k=5, metadata_filter=None):
        if metadata_filter:
            return [(owner_doc, 0), (other_doc, 0.1)]
        return [(enterprise_doc, 0.2), (other_doc, 0.3)]

    with patch(
        "app.retrieval.vector_store.async_similarity_search",
        new=AsyncMock(side_effect=fake_search),
    ), patch(
        "app.retrieval.vector_store.async_get_document_chunks",
        new=AsyncMock(return_value=[]),
    ):
        results = await retriever._chroma_search(
            "测试",
            top_k=5,
            knowledge_scope=["enterprise", "personal"],
            owner_open_id="ou_test123",
        )

    titles = [item.title for item in results]
    assert "enterprise.txt" in titles
    assert "mine.txt" in titles
    assert "other.txt" not in titles


@pytest.mark.asyncio
async def test_chroma_search_merges_opensearch_lexical_results():
    retriever = Retriever()
    lexical_doc = MagicMock(
        page_content="大语言模型的训练通常分为预训练和微调两个阶段",
        metadata={
            "source": "llm.txt",
            "knowledge_base_type": "enterprise",
            "document_id": "doc_os",
            "chunk_index": 0,
        },
    )

    with patch(
        "app.retrieval.vector_store.async_similarity_search",
        new=AsyncMock(return_value=[]),
    ), patch(
        "app.retrieval.opensearch_store.async_opensearch_search_with_status",
        new=AsyncMock(return_value=OpenSearchSearchResult(results=[(lexical_doc, 0.01)], available=True, succeeded=True)),
    ), patch(
        "app.retrieval.vector_store.async_get_document_chunks",
        new=AsyncMock(return_value=[]),
    ):
        results = await retriever._chroma_search(
            "大模型训练",
            top_k=1,
            knowledge_scope=["enterprise"],
        )

    assert len(results) == 1
    assert results[0].title == "llm.txt"
    assert "大语言模型的训练" in results[0].snippet


@pytest.mark.asyncio
async def test_combined_scope_keeps_personal_result_when_enterprise_scores_dominate():
    retriever = Retriever()
    enterprise_docs = [
        MagicMock(
            page_content=f"绘制一条正弦曲线 企业片段 {index}",
            metadata={"source": f"enterprise_{index}.txt", "knowledge_base_type": "enterprise"},
        )
        for index in range(6)
    ]
    personal_doc = MagicMock(
        page_content="Python实现快速排序算法\ndef quick_sort(arr):\n    return arr",
        metadata={"source": "personal_quick_sort.txt", "knowledge_base_type": "personal", "owner_open_id": "ou_test123"},
    )

    async def fake_search(_query, k=5, metadata_filter=None):
        if metadata_filter:
            return [(personal_doc, 0.9)]
        return [(doc, index * 0.01) for index, doc in enumerate(enterprise_docs)]

    with patch(
        "app.retrieval.vector_store.async_similarity_search",
        new=AsyncMock(side_effect=fake_search),
    ), patch(
        "app.retrieval.vector_store.async_get_document_chunks",
        new=AsyncMock(return_value=[]),
    ):
        results = await retriever._chroma_search(
            "绘制一条正弦曲线 Python实现快速排序算法",
            top_k=5,
            knowledge_scope=["enterprise", "personal"],
            owner_open_id="ou_test123",
        )

    titles = [item.title for item in results]
    assert "personal_quick_sort.txt" in titles
    assert any(title.startswith("enterprise_") for title in titles)


@pytest.mark.asyncio
async def test_combined_scope_searches_quoted_targets_separately():
    retriever = Retriever()
    animation_doc = MagicMock(
        page_content="制作一个圆点沿曲线运动的动画，并时刻显示圆点的坐标位置\n\ndef update(frame):\n    return point, coord_text",
        metadata={"source": "enterprise_animation.txt", "knowledge_base_type": "enterprise"},
    )
    quick_sort_doc = MagicMock(
        page_content="Python实现快速排序算法\ndef quick_sort(arr):\n    return quick_sort(left) + middle + quick_sort(right)",
        metadata={"source": "personal_quick_sort.txt", "knowledge_base_type": "personal", "owner_open_id": "ou_test123"},
    )
    irrelevant_personal_doc = MagicMock(
        page_content="个人库无关内容",
        metadata={"source": "personal_other.txt", "knowledge_base_type": "personal", "owner_open_id": "ou_test123"},
    )

    async def fake_search(query, k=5, metadata_filter=None):
        if "快速排序" in query and metadata_filter:
            return [(quick_sort_doc, 0)]
        if "圆点沿曲线运动" in query and metadata_filter is None:
            return [(animation_doc, 0)]
        if metadata_filter:
            return [(irrelevant_personal_doc, 1.5)]
        return []

    with patch(
        "app.retrieval.vector_store.async_similarity_search",
        new=AsyncMock(side_effect=fake_search),
    ), patch(
        "app.retrieval.vector_store.async_get_document_chunks",
        new=AsyncMock(return_value=[]),
    ):
        results = await retriever._chroma_search(
            "企业知识库和个人知识库结合检索，找出“制作一个圆点沿曲线运动的动画，并时刻显示圆点的坐标位置”和“Python实现快速排序算法”的python代码片段",
            top_k=5,
            knowledge_scope=["enterprise", "personal"],
            owner_open_id="ou_test123",
        )

    titles = [item.title for item in results]
    assert "enterprise_animation.txt" in titles
    assert "personal_quick_sort.txt" in titles


@pytest.mark.asyncio
async def test_combined_scope_preserves_one_result_per_quoted_target_before_top_k_cutoff():
    retriever = Retriever()
    animation_docs = [
        MagicMock(
            page_content=f"制作一个圆点沿曲线运动的动画 片段 {index}",
            metadata={"source": f"animation_{index}.txt", "knowledge_base_type": "enterprise"},
        )
        for index in range(5)
    ]
    quick_sort_doc = MagicMock(
        page_content="Python实现快速排序算法\ndef quick_sort(arr):\n    return arr",
        metadata={"source": "personal_quick_sort.txt", "knowledge_base_type": "personal", "owner_open_id": "ou_request_user"},
    )

    async def fake_search(query, k=5, metadata_filter=None):
        if "圆点沿曲线运动" in query and metadata_filter is None:
            return [(doc, index * 0.01) for index, doc in enumerate(animation_docs)]
        if "快速排序" in query and metadata_filter:
            return [(quick_sort_doc, 0.8)]
        return []

    with patch(
        "app.retrieval.vector_store.async_similarity_search",
        new=AsyncMock(side_effect=fake_search),
    ), patch(
        "app.retrieval.vector_store.async_get_document_chunks",
        new=AsyncMock(return_value=[]),
    ):
        results = await retriever._chroma_search(
            "找出“制作一个圆点沿曲线运动的动画”和“Python实现快速排序算法”的python代码片段",
            top_k=2,
            knowledge_scope=["enterprise", "personal"],
            owner_open_id="ou_request_user",
        )

    titles = [item.title for item in results]
    assert len(results) == 2
    assert any(title.startswith("animation_") for title in titles)
    assert "personal_quick_sort.txt" in titles


@pytest.mark.asyncio
async def test_quoted_target_search_runs_subquestions_in_parallel():
    retriever = Retriever()
    active = 0
    max_active = 0
    first_doc = MagicMock(
        page_content="创意写作与描述性文本",
        metadata={"source": "creative.txt", "knowledge_base_type": "enterprise"},
    )
    second_doc = MagicMock(
        page_content="大同思想的历史影响",
        metadata={"source": "datong.txt", "knowledge_base_type": "enterprise"},
    )

    async def fake_search(query, k=5, metadata_filter=None):
        nonlocal active, max_active
        active += 1
        max_active = max(max_active, active)
        await asyncio.sleep(0.05)
        active -= 1
        if "创意写作" in query:
            return [(first_doc, 0)]
        if "大同思想" in query:
            return [(second_doc, 0)]
        return []

    with patch(
        "app.retrieval.vector_store.async_similarity_search",
        new=AsyncMock(side_effect=fake_search),
    ), patch(
        "app.retrieval.opensearch_store.async_opensearch_search_with_status",
        new=AsyncMock(return_value=OpenSearchSearchResult(results=[], available=True, succeeded=True)),
    ), patch(
        "app.retrieval.vector_store.async_get_document_chunks",
        new=AsyncMock(return_value=[]),
    ):
        results = await retriever._chroma_search(
            "查找“创意写作与描述性文本”和“大同思想的历史影响”的相关文本",
            top_k=2,
            knowledge_scope=["enterprise"],
        )

    assert max_active > 1
    titles = [item.title for item in results]
    assert "creative.txt" in titles
    assert "datong.txt" in titles


@pytest.mark.asyncio
async def test_enterprise_keyword_search_recalls_weather_xlsx_chunk():
    retriever = Retriever()
    noise_doc = MagicMock(
        page_content="Matplotlib绘图代码和人工智能伦理问题",
        metadata={"source": "noise.txt", "knowledge_base_type": "enterprise"},
    )
    weather_doc = MagicMock(
        page_content="郑丽华\t20125081042\t基于推荐算法的电子商城设计与实现\t基于BP神经网络的碳排放数据分析系统的设计与实现\t基于机器学习的天气数据分析与预测系统的实现",
        metadata={
            "source": "weather.xlsx",
            "knowledge_base_type": "enterprise",
            "document_id": "xlsx1",
            "chunk_index": 0,
        },
    )

    async def fake_similarity_search(_query, k=5, metadata_filter=None):
        return [(noise_doc, 0)]

    async def fake_keyword_search(terms, k=5, metadata_filter=None):
        assert "天气" in terms
        return [(weather_doc, -10)]

    with patch(
        "app.retrieval.vector_store.async_similarity_search",
        new=AsyncMock(side_effect=fake_similarity_search),
    ), patch(
        "app.retrieval.vector_store.async_keyword_search",
        new=AsyncMock(side_effect=fake_keyword_search),
    ), patch(
        "app.retrieval.vector_store.async_get_document_chunks",
        new=AsyncMock(return_value=[]),
    ):
        results = await retriever._chroma_search(
            "查找关于天气的历年毕业选题",
            top_k=3,
            knowledge_scope=["enterprise"],
        )

    assert results[0].title == "weather.xlsx"
    assert "天气数据分析" in results[0].snippet
