from pathlib import Path
from unittest.mock import AsyncMock, patch

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.documents.adapters.docx_http import HttpDocxBackend
from app.documents.adapters.xlsx_http import HttpXlsxBackend
from services.docx_mcp.operations import (
    clear_table_cell,
    read_structure,
    read_table_cell,
    replace_table_cell,
)
from services.docx_mcp.main import app as docx_app
from services.mcp_common.path_security import validate_file_path


def test_path_map_translates_windows_host_path(monkeypatch, tmp_path):
    upload_dir = tmp_path / "uploads"
    upload_dir.mkdir()
    target = upload_dir / "sample.docx"
    target.write_bytes(b"docx")

    monkeypatch.setenv("MCP_ALLOWED_DIR", str(upload_dir))
    monkeypatch.setenv("MCP_PATH_MAP", r"D:\LLM\Unified_API_service\data\uploads=" + str(upload_dir))

    resolved, err = validate_file_path(r"D:\LLM\Unified_API_service\data\uploads\sample.docx")

    assert err is None
    assert resolved == target.resolve()


def test_path_map_does_not_match_partial_prefix(monkeypatch, tmp_path):
    upload_dir = tmp_path / "uploads"
    upload_dir.mkdir()
    outside_dir = tmp_path / "uploads-other"
    outside_dir.mkdir()
    target = outside_dir / "sample.docx"
    target.write_bytes(b"docx")

    monkeypatch.setenv("MCP_ALLOWED_DIR", str(upload_dir))
    monkeypatch.setenv("MCP_PATH_MAP", f"{tmp_path / 'uploads'}={upload_dir}")

    _, err = validate_file_path(str(target))

    assert err is not None


def test_docx_mcp_requires_api_key_when_configured(monkeypatch):
    monkeypatch.setenv("MCP_API_KEY", "secret")
    client = TestClient(docx_app)

    response = client.post("/docx/read_structure", json={"file_path": "/tmp/sample.docx"})

    assert response.status_code == 401


def test_docx_mcp_rejects_wrong_api_key(monkeypatch):
    monkeypatch.setenv("MCP_API_KEY", "secret")
    client = TestClient(docx_app)

    response = client.post(
        "/docx/read_structure",
        headers={"X-MCP-API-Key": "wrong"},
        json={"file_path": "/tmp/sample.docx"},
    )

    assert response.status_code == 403


@pytest.mark.asyncio
async def test_docx_http_backend_sends_api_key_header():
    backend = HttpDocxBackend("http://localhost:9000", timeout=5.0, api_key="secret")
    mock_resp = httpx.Response(200, json={
        "success": True,
        "data": {"type": "docx", "paragraph_count": 3},
    })
    mock_post = AsyncMock(return_value=mock_resp)

    with patch("httpx.AsyncClient.post", new=mock_post):
        await backend.read_structure("/tmp/test.docx")

    assert mock_post.call_args.kwargs["headers"] == {"X-MCP-API-Key": "secret"}


@pytest.mark.asyncio
async def test_xlsx_http_backend_sends_api_key_header():
    backend = HttpXlsxBackend("http://localhost:9001", timeout=5.0, api_key="secret")
    mock_resp = httpx.Response(200, json={
        "success": True,
        "data": {"value": "ok"},
    })
    mock_post = AsyncMock(return_value=mock_resp)

    with patch("httpx.AsyncClient.post", new=mock_post):
        await backend.read_cell("/tmp/test.xlsx", "Sheet1", "A1")

    assert mock_post.call_args.kwargs["headers"] == {"X-MCP-API-Key": "secret"}


def test_docx_read_structure_includes_table_cell_text(tmp_path):
    docx_path = tmp_path / "table.docx"
    doc = Document()
    doc.add_paragraph("普通段落")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "标题"
    table.cell(0, 1).text = "说明"
    table.cell(1, 0).text = (
        "实验目的：\n"
        "（1）掌握 Pandas 读取数据及 Matplotlib 绘制图表的方法；\n"
        "（2）学会生成随机数据并绘制直方图。"
    )
    table.cell(1, 1).text = "实验内容"
    doc.save(docx_path)

    structure = read_structure(str(docx_path))

    assert structure["paragraphs"][0]["style"]
    cells = structure["tables"][0]["cells"]
    purpose_cell = next(cell for cell in cells if "实验目的" in cell["text"])
    assert purpose_cell["row"] == 1
    assert purpose_cell["col"] == 0
    assert "\n" in purpose_cell["text"]


def test_docx_read_structure_dedupes_merged_cells(tmp_path):
    docx_path = tmp_path / "merged.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "合并单元格内容"
    doc.save(docx_path)

    structure = read_structure(str(docx_path))

    cells = structure["tables"][0]["cells"]
    assert [cell["text"] for cell in cells].count("合并单元格内容") == 1


def _create_docx_with_table(path: Path) -> None:
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "标题"
    table.cell(0, 1).text = "说明"
    table.cell(1, 0).text = "实验目的：掌握 Pandas 读取数据及 Matplotlib 绘图方法"
    table.cell(1, 1).text = "实验内容"
    doc.save(path)


def test_docx_table_cell_operations_read_replace_clear(tmp_path):
    docx_path = tmp_path / "table_ops.docx"
    _create_docx_with_table(docx_path)

    text = read_table_cell(str(docx_path), 0, 1, 0)
    assert "实验目的" in text

    replace_table_cell(str(docx_path), 0, 1, 0, "新的实验目的")
    assert read_table_cell(str(docx_path), 0, 1, 0) == "新的实验目的"

    clear_table_cell(str(docx_path), 0, 1, 0)
    assert read_table_cell(str(docx_path), 0, 1, 0) == ""


def test_docx_table_cell_operations_reject_out_of_range(tmp_path):
    docx_path = tmp_path / "table_out_of_range.docx"
    _create_docx_with_table(docx_path)

    with pytest.raises(IndexError):
        read_table_cell(str(docx_path), 1, 0, 0)
    with pytest.raises(IndexError):
        replace_table_cell(str(docx_path), 0, 9, 0, "x")
    with pytest.raises(IndexError):
        clear_table_cell(str(docx_path), 0, 0, 9)


@pytest.mark.asyncio
async def test_docx_http_backend_table_cell_endpoints():
    backend = HttpDocxBackend("http://localhost:9000", timeout=5.0, api_key="secret")
    mock_post = AsyncMock(side_effect=[
        httpx.Response(200, json={"success": True, "data": {"text": "实验目的"}}),
        httpx.Response(200, json={"success": True, "data": {}}),
        httpx.Response(200, json={"success": True, "data": {}}),
    ])

    with patch("httpx.AsyncClient.post", new=mock_post):
        text = await backend.read_table_cell("/tmp/test.docx", 0, 1, 0)
        await backend.replace_table_cell("/tmp/test.docx", 0, 1, 0, "新内容")
        await backend.clear_table_cell("/tmp/test.docx", 0, 1, 0)

    assert text == "实验目的"
    called_paths = [call.args[0] for call in mock_post.call_args_list]
    assert called_paths == [
        "http://localhost:9000/docx/read_table_cell",
        "http://localhost:9000/docx/replace_table_cell",
        "http://localhost:9000/docx/clear_table_cell",
    ]
    replace_payload = mock_post.call_args_list[1].kwargs["json"]
    assert replace_payload == {
        "file_path": "/tmp/test.docx",
        "table_index": 0,
        "row": 1,
        "col": 0,
        "new_text": "新内容",
    }
