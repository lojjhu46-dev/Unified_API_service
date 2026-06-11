"""DOCX 文档操作

基于 python-docx 的同步操作实现。
"""

from __future__ import annotations

import shutil
from typing import Any

from docx import Document


def _paragraph_style_name(paragraph) -> str:
    if paragraph.style and paragraph.style.name:
        return paragraph.style.name
    return ""


def _truncate_text(text: str, max_length: int) -> str:
    if len(text) > max_length:
        return text[:max_length] + "..."
    return text


def _get_table_cell(doc: Document, table_index: int, row: int, col: int):
    if table_index < 0 or table_index >= len(doc.tables):
        raise IndexError(f"表格索引 {table_index} 超出范围（共 {len(doc.tables)} 个表格）")
    table = doc.tables[table_index]
    if row < 0 or row >= len(table.rows):
        raise IndexError(f"表格 {table_index} 行索引 {row} 超出范围（共 {len(table.rows)} 行）")
    row_cells = table.rows[row].cells
    if col < 0 or col >= len(row_cells):
        raise IndexError(f"表格 {table_index} 列索引 {col} 超出范围（该行共 {len(row_cells)} 列）")
    return row_cells[col]


def read_structure(file_path: str) -> dict[str, Any]:
    """读取 DOCX 文档结构"""
    doc = Document(file_path)
    paragraphs = doc.paragraphs

    headings = []
    paragraphs_info = []
    max_paragraphs = 200  # 最多返回前 200 段
    max_text_length = 500  # 每段最多 500 字

    for i, p in enumerate(paragraphs[:max_paragraphs]):
        text = _truncate_text(p.text, max_text_length)
        style_name = _paragraph_style_name(p)

        paragraphs_info.append({
            "index": i,
            "text": text,
            "style": style_name,
        })

        if style_name.startswith("Heading"):
            headings.append(p.text)

    tables_info = []
    max_tables = 20
    max_cells_per_table = 500
    max_cell_text_length = 1000
    for i, table in enumerate(doc.tables):
        cells = []
        seen_cell_elements = set()
        for row_index, row in enumerate(table.rows):
            for col_index, cell in enumerate(row.cells):
                cell_element = cell._tc
                if cell_element in seen_cell_elements:
                    continue
                seen_cell_elements.add(cell_element)
                text = _truncate_text(cell.text, max_cell_text_length)
                cells.append({
                    "row": row_index,
                    "col": col_index,
                    "text": text,
                    "merged": False,
                })
                if len(cells) >= max_cells_per_table:
                    break
            if len(cells) >= max_cells_per_table:
                break
        tables_info.append({
            "index": i,
            "rows": len(table.rows),
            "cols": len(table.columns),
            "cells": cells,
            "truncated": len(cells) >= max_cells_per_table,
        })
        if len(tables_info) >= max_tables:
            break

    return {
        "type": "docx",
        "paragraph_count": len(paragraphs),
        "paragraphs": paragraphs_info,
        "headings": headings,
        "tables": tables_info,
        "tables_truncated": len(doc.tables) > max_tables,
    }


def read_paragraph(file_path: str, index: int) -> str:
    """读取指定段落的文本"""
    doc = Document(file_path)
    paragraphs = doc.paragraphs

    if index < 0 or index >= len(paragraphs):
        raise IndexError(f"段落索引 {index} 超出范围（共 {len(paragraphs)} 段）")

    return paragraphs[index].text


def replace_paragraph(file_path: str, index: int, new_text: str) -> None:
    """替换指定段落的文本"""
    doc = Document(file_path)
    paragraphs = doc.paragraphs

    if index < 0 or index >= len(paragraphs):
        raise IndexError(f"段落索引 {index} 超出范围（共 {len(paragraphs)} 段）")

    # 清除原有内容并添加新文本，保留段落样式
    paragraph = paragraphs[index]
    paragraph.clear()
    paragraph.add_run(new_text)

    doc.save(file_path)


def delete_paragraph(file_path: str, index: int) -> None:
    """删除指定段落"""
    doc = Document(file_path)
    paragraphs = doc.paragraphs

    if index < 0 or index >= len(paragraphs):
        raise IndexError(f"段落索引 {index} 超出范围（共 {len(paragraphs)} 段）")

    # 从 XML 树中移除段落元素
    paragraph = paragraphs[index]
    parent = paragraph._element.getparent()
    if parent is not None:
        parent.remove(paragraph._element)

    doc.save(file_path)


def append_paragraph(file_path: str, text: str) -> None:
    """在文档末尾追加段落"""
    doc = Document(file_path)
    doc.add_paragraph(text)
    doc.save(file_path)


def read_table_cell(file_path: str, table_index: int, row: int, col: int) -> str:
    """读取指定表格单元格文本"""
    doc = Document(file_path)
    cell = _get_table_cell(doc, table_index, row, col)
    return cell.text


def replace_table_cell(file_path: str, table_index: int, row: int, col: int, new_text: str) -> None:
    """替换指定表格单元格文本"""
    doc = Document(file_path)
    cell = _get_table_cell(doc, table_index, row, col)
    cell.text = new_text
    doc.save(file_path)


def clear_table_cell(file_path: str, table_index: int, row: int, col: int) -> None:
    """清空指定表格单元格文本"""
    replace_table_cell(file_path, table_index, row, col, "")


def save_copy(file_path: str, output_path: str) -> str:
    """保存文档副本到指定路径"""
    shutil.copy2(file_path, output_path)
    return output_path
