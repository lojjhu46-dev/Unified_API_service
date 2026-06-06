"""XLSX 文档操作

基于 openpyxl 的同步操作实现。
"""

from __future__ import annotations

import shutil
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def _column_name(index: int) -> str:
    """将 0-based 列索引转换为 Excel 列名：0→A, 25→Z, 26→AA, 27→AB, ..."""
    if index < 0:
        raise ValueError(f"列索引必须 >= 0，当前为 {index}")
    name = ""
    while True:
        name = chr(ord("A") + index % 26) + name
        index = index // 26 - 1
        if index < 0:
            break
    return name


def read_structure(file_path: str) -> dict[str, Any]:
    """读取 XLSX 工作簿结构"""
    wb = load_workbook(file_path, read_only=True)
    sheets_info = {}

    # 顶层字段默认取第一个 sheet
    first_sheet_headers = []
    first_sheet_row_count = 0

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        # 读取表头（第一行）
        headers = []
        for cell in ws[1]:
            headers.append(str(cell.value) if cell.value is not None else "")

        row_count = ws.max_row or 0

        sheets_info[sheet_name] = {
            "headers": headers,
            "row_count": row_count,
        }

        # 第一个 sheet 的数据作为顶层默认值
        if not first_sheet_headers:
            first_sheet_headers = headers
            first_sheet_row_count = row_count

    wb.close()

    return {
        "type": "xlsx",
        "sheets": list(wb.sheetnames),
        "headers": first_sheet_headers,
        "row_count": first_sheet_row_count,
        "sheets_info": sheets_info,
    }


def read_cell(file_path: str, sheet: str, cell: str) -> str:
    """读取指定单元格的值"""
    wb = load_workbook(file_path, read_only=True)

    if sheet not in wb.sheetnames:
        wb.close()
        raise KeyError(f"工作表不存在: {sheet}")

    ws = wb[sheet]
    value = ws[cell.upper()].value
    wb.close()

    return str(value) if value is not None else ""


def modify_cell(file_path: str, sheet: str, cell: str, value: str) -> None:
    """修改指定单元格的值"""
    wb = load_workbook(file_path)

    if sheet not in wb.sheetnames:
        wb.close()
        raise KeyError(f"工作表不存在: {sheet}")

    ws = wb[sheet]
    ws[cell.upper()] = value
    wb.save(file_path)
    wb.close()


def append_row(file_path: str, sheet: str, values: list[str]) -> None:
    """在指定工作表末尾追加一行"""
    wb = load_workbook(file_path)

    if sheet not in wb.sheetnames:
        wb.close()
        raise KeyError(f"工作表不存在: {sheet}")

    ws = wb[sheet]
    ws.append(values)
    wb.save(file_path)
    wb.close()


def delete_row(file_path: str, sheet: str, row: int) -> None:
    """删除指定行（1-based）"""
    if row < 1:
        raise ValueError(f"行号必须 >= 1，当前为 {row}")

    wb = load_workbook(file_path)

    if sheet not in wb.sheetnames:
        wb.close()
        raise KeyError(f"工作表不存在: {sheet}")

    ws = wb[sheet]
    ws.delete_rows(row, 1)
    wb.save(file_path)
    wb.close()


def save_copy(file_path: str, output_path: str) -> str:
    """保存文档副本到指定路径"""
    shutil.copy2(file_path, output_path)
    return output_path
