"""文档摄取模块"""

import os
import re
import uuid
from pathlib import Path
from app.config import settings
from app.observability.logging import get_logger

logger = get_logger(__name__)

_INVALID_FILENAME_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
_ALLOWED_EXTENSIONS = {".pdf", ".txt", ".docx", ".xlsx"}


def load_document(file_path: str) -> list:
    """加载文档"""
    try:
        from langchain_community.document_loaders import Docx2txtLoader, PyPDFLoader, TextLoader

        ext = os.path.splitext(file_path)[1].lower()
        if ext == ".pdf":
            loader = PyPDFLoader(file_path)
        elif ext == ".txt":
            loader = TextLoader(file_path, encoding="utf-8")
        elif ext == ".docx":
            loader = Docx2txtLoader(file_path)
        elif ext == ".xlsx":
            return load_xlsx_document(file_path)
        else:
            raise ValueError(f"不支持的文件格式: {ext}")

        return loader.load()
    except Exception as e:
        logger.error(f"加载文档失败: {e}")
        raise


def load_xlsx_document(file_path: str) -> list:
    """将XLSX按工作表抽取为文本Document。"""
    from langchain_core.documents import Document
    from openpyxl import load_workbook

    workbook = load_workbook(file_path, data_only=True, read_only=True)
    documents = []
    try:
        for sheet in workbook.worksheets:
            lines = []
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(value.strip() for value in values):
                    lines.append("\t".join(values).rstrip())
            content = "\n".join(lines).strip()
            if content:
                documents.append(Document(
                    page_content=content,
                    metadata={"source": file_path, "sheet": sheet.title},
                ))
    finally:
        workbook.close()
    return documents


def split_documents(documents: list) -> list:
    """切分文档"""
    try:
        from langchain_text_splitters import RecursiveCharacterTextSplitter

        splitter = RecursiveCharacterTextSplitter(
            chunk_size=settings.chunk_size,
            chunk_overlap=settings.chunk_overlap,
            separators=["\n\n", "\n", "。", "！", "？", "，", " ", ""],
        )
        return splitter.split_documents(documents)
    except Exception as e:
        logger.error(f"切分文档失败: {e}")
        raise


def ingest_file(
    file_path: str,
    original_filename: str | None = None,
    knowledge_base_type: str = "enterprise",
    owner_open_id: str | None = None,
    owner_user_id: str | None = None,
    tenant_id: str | None = None,
    document_id: str | None = None,
    chat_id: str | None = None,
    channel: str = "api",
) -> dict:
    """摄取文件到向量存储"""
    from app.retrieval.document_registry import document_registry
    from app.retrieval.vector_store import add_documents
    from app.retrieval.opensearch_store import index_documents

    document_id = document_id or str(uuid.uuid4())[:12]
    tenant_id = tenant_id or settings.default_tenant_id
    knowledge_base_type = (knowledge_base_type or "enterprise").strip().lower()
    if knowledge_base_type == "personal":
        owner_user_id = owner_user_id or owner_open_id or ""
        owner_open_id = owner_open_id or owner_user_id or ""
    else:
        owner_user_id = owner_user_id or ""
        owner_open_id = owner_open_id or ""
    stored_filename = os.path.basename(file_path)
    safe_original_filename = sanitize_filename(original_filename or stored_filename)

    document_registry.create_processing(
        document_id=document_id,
        tenant_id=tenant_id,
        knowledge_base_type=knowledge_base_type,
        owner_user_id=owner_user_id,
        owner_open_id=owner_open_id,
        original_filename=safe_original_filename,
        stored_filename=stored_filename,
        stored_path=str(Path(file_path).resolve()),
        channel=channel,
        chat_id=chat_id or "",
    )

    try:
        docs = load_document(file_path)
        chunks = split_documents(docs)
        for index, chunk in enumerate(chunks):
            # 这些元数据会写入 Chroma，后续可用于来源追踪、权限过滤和文档删除。
            metadata = chunk.metadata if isinstance(getattr(chunk, "metadata", None), dict) else {}
            chunk.metadata = {
                **metadata,
                "tenant_id": tenant_id,
                "document_id": document_id,
                "chunk_id": f"{document_id}:{index}",
                "original_filename": safe_original_filename,
                "stored_filename": stored_filename,
                "chunk_index": index,
                "knowledge_base_type": knowledge_base_type,
                "owner_user_id": owner_user_id,
                "owner_open_id": owner_open_id,
                "chat_id": chat_id or "",
                "channel": channel,
            }
        chunk_count = add_documents(chunks)
        opensearch_count = index_documents(chunks)
        if opensearch_count:
            logger.info(f"OpenSearch index updated: {opensearch_count}个切块")

        document_registry.mark_ready(document_id, chunk_count)
        logger.info(f"文档摄取完成: {safe_original_filename}, {chunk_count}个切块")
    except Exception as e:
        document_registry.mark_failed(document_id, str(e))
        raise

    return {
        "document_id": document_id,
        "filename": safe_original_filename,
        "chunks": chunk_count,
    }


def validate_file_extension(filename: str) -> bool:
    """验证文件扩展名"""
    ext = os.path.splitext(filename or "")[1].lower()
    return ext in _ALLOWED_EXTENSIONS


def sanitize_filename(filename: str | None) -> str:
    """清理上传文件名，防止路径穿越和非法字符进入保存路径。"""
    raw_name = filename or ""
    # UploadFile.filename 可能携带 Windows 或 Unix 路径分隔符，这里统一只保留 basename。
    base_name = raw_name.replace("\\", "/").split("/")[-1]
    ext = os.path.splitext(base_name)[1].lower()
    if ext not in _ALLOWED_EXTENSIONS:
        ext = ".txt"

    stem = os.path.splitext(base_name)[0]
    stem = _INVALID_FILENAME_CHARS.sub("_", stem).strip(" .")
    if not stem:
        stem = "upload"
    return f"{stem}{ext}"


def save_uploaded_file(
    content: bytes,
    filename: str,
    upload_dir: str | None = None,
    *,
    tenant_id: str | None = None,
    owner_user_id: str | None = None,
    knowledge_base_type: str | None = None,
    document_id: str | None = None,
) -> str:
    """保存上传的文件"""
    target_dir = Path(upload_dir or settings.upload_dir)
    if knowledge_base_type:
        kb_type = (knowledge_base_type or "enterprise").strip().lower()
        tenant = sanitize_path_segment(tenant_id or settings.default_tenant_id)
        if kb_type == "personal":
            owner = sanitize_path_segment(owner_user_id or "unknown")
            target_dir = target_dir / "personal" / tenant / owner
        else:
            target_dir = target_dir / "enterprise" / tenant
    target_dir = target_dir.resolve()
    target_dir.mkdir(parents=True, exist_ok=True)

    safe_filename = sanitize_filename(filename)
    prefix = document_id or str(uuid.uuid4())
    file_path = (target_dir / f"{prefix}_{safe_filename}").resolve()
    if target_dir not in file_path.parents:
        raise ValueError("上传文件路径非法")

    with open(file_path, "wb") as f:
        f.write(content)
    return str(file_path)


def sanitize_path_segment(value: str | None) -> str:
    """清理租户和用户 ID 路径片段。"""
    segment = re.sub(r"[^A-Za-z0-9_.-]+", "_", value or "").strip("._")
    return segment or "default"
